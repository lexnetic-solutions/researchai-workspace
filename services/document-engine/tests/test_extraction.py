"""End-to-end extraction tests against generated fixture files."""

from __future__ import annotations

from pathlib import Path

from researchai_document_engine.parsers import get_parser


def _write(tmp: Path, name: str, data: bytes) -> str:
    f = tmp / name
    f.write_bytes(data)
    return str(f)


def test_markdown_headings_sections_offsets(tmp_path: Path) -> None:
    content = (
        "# Intro\n\nFirst paragraph text.\n\n"
        "## Methods\n\nSecond paragraph under methods.\n"
    )
    path = _write(tmp_path, "sample.md", content.encode())
    res = get_parser("md")(path)
    assert res.ok
    assert res.title == "sample"
    headings = [s.heading for s in res.sections]
    assert "Intro" in headings and "Methods" in headings
    methods = next(s for s in res.sections if s.heading == "Methods")
    assert methods.level == 2
    # Block offsets must land inside the source text.
    for b in res.blocks:
        assert res.blocks is not None
        assert content[b.start_offset : b.end_offset] == b.text


def test_txt_paragraphs_roundtrip(tmp_path: Path) -> None:
    content = "Alpha paragraph.\n\nBeta paragraph.\n"
    path = _write(tmp_path, "notes.txt", content.encode())
    res = get_parser("txt")(path)
    assert res.ok
    texts = [b.text for b in res.blocks]
    assert any("Alpha" in t for t in texts)
    assert any("Beta" in t for t in texts)
    for b in res.blocks:
        assert content[b.start_offset : b.end_offset] == b.text


def test_docx_headings_and_blocks(tmp_path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Study Design", level=1)
    document.add_paragraph("Participants were recruited locally.")
    document.add_heading("Outcomes", level=2)
    document.add_paragraph("Primary outcome improved.")
    path = str(tmp_path / "study.docx")
    document.save(path)

    res = get_parser("docx")(path)
    assert res.ok, res.error
    headings = [s.heading for s in res.sections]
    assert "Study Design" in headings and "Outcomes" in headings
    outcomes = next(s for s in res.sections if s.heading == "Outcomes")
    assert outcomes.level == 2
    assert any("Participants were recruited locally." == b.text for b in res.blocks)


def test_pdf_pages_blocks_title(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_metadata({"/Title": "Fixture Paper"})
    path = str(tmp_path / "paper.pdf")
    with open(path, "wb") as fh:
        writer.write(fh)

    res = get_parser("pdf")(path)
    assert res.ok, res.error
    assert res.page_count == 1
    assert res.title == "Fixture Paper"
    assert len(res.sections) == 1
    assert res.sections[0].page_start == 1


def test_corrupt_pdf_returns_honest_error(tmp_path: Path) -> None:
    path = _write(tmp_path, "broken.pdf", b"%PDF-1.4 this is not really a pdf")
    res = get_parser("pdf")(path)
    assert not res.ok
    assert res.error


def test_html_headings_extracted(tmp_path: Path) -> None:
    html = b"<html><head><style>x{}</style></head><body><h1>Title</h1><p>Body text</p></body></html>"
    path = _write(tmp_path, "page.html", html)
    res = get_parser("html")(path)
    assert res.ok, res.error
    headings = [s.heading for s in res.sections]
    assert "Title" in headings
