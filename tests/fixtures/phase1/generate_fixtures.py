"""Generate binary fixtures for the live ingestion test.

Run from the engine venv (has python-docx + pypdf):
    cd services/document-engine && uv run python ../../tests/fixtures/phase1/generate_fixtures.py
"""

from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).parent


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Study Design", level=1)
    document.add_paragraph("Participants were recruited from the local university population.")
    document.add_heading("Outcomes", level=2)
    document.add_paragraph("The primary outcome improved across all measured conditions.")
    document.save(path)
    print(f"wrote {path}")


def make_pdf(path: Path) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_metadata({"/Title": "Fixture Paper"})
    with open(path, "wb") as fh:
        writer.write(fh)
    print(f"wrote {path}")


if __name__ == "__main__":
    which = sys.argv[1] if len(sys.argv) > 1 else "all"
    if which in ("all", "docx"):
        make_docx(HERE / "sample.docx")
    if which in ("all", "pdf"):
        make_pdf(HERE / "sample.pdf")
